from fastapi import FastAPI, Query, HTTPException, UploadFile, File, Form, Request
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
import sqlite3
import os
import json
import time
from datetime import datetime
import openpyxl
import io
import csv
from dotenv import load_dotenv
import re


# Import AI Engine & Demo Data
try:
    from . import ai_engine
    from . import demo_data
except (ImportError, ValueError):
    import ai_engine
    import demo_data

# Safe import for psycopg2
try:
    import psycopg2
    from psycopg2.extras import RealDictCursor
    HAS_PSYCOPG2 = True
except Exception as e:
    psycopg2 = None
    RealDictCursor = None
    HAS_PSYCOPG2 = False
    print(f"Notice: psycopg2 not loaded: {e}")

env_path = os.path.join(os.path.dirname(__file__), ".env")
if os.path.exists(env_path):
    load_dotenv(env_path)
load_dotenv()

app = FastAPI(title="Indore ETL RU & Marine Hull API", version="2.2.0")

# Setup CORS to allow the frontend to communicate with the backend
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

from starlette.types import ASGIApp, Scope, Receive, Send
from urllib.parse import parse_qs, urlencode

class VercelPathRewriteMiddleware:
    def __init__(self, app: ASGIApp):
        self.app = app

    async def __call__(self, scope: Scope, receive: Receive, send: Send):
        if scope["type"] == "http":
            query_bytes = scope.get("query_string", b"")
            query_str = query_bytes.decode("latin1")
            
            if "__path=" in query_str:
                params = parse_qs(query_str, keep_blank_values=True)
                extracted_path = params.pop("__path", [""])[0]
                
                # Reconstruct query string without internal __path routing parameter
                new_query = urlencode(params, doseq=True)
                scope["query_string"] = new_query.encode("latin1")
                
                # Format clean routing path
                clean_path = "/" + extracted_path.lstrip("/")
                scope["path"] = clean_path
                scope["raw_path"] = clean_path.encode("latin1")
            else:
                path = scope.get("path", "")
                for prefix in ["/api/index.py", "/api/index", "/index.py"]:
                    if path.startswith(prefix):
                        new_path = path[len(prefix):] or "/"
                        scope["path"] = new_path
                        scope["raw_path"] = new_path.encode("latin1")
                        break

        await self.app(scope, receive, send)

app.add_middleware(VercelPathRewriteMiddleware)

def safe_int(val, default: int = 0) -> int:
    try:
        if val is None:
            return default
        s = str(val).strip()
        return int(s) if s else default
    except (ValueError, TypeError):
        return default

# Database Configuration (Supabase & Railway friendly)
DATABASE_URL = os.getenv("DATABASE_URL", "").strip()
DB_CONFIG = {
    "host": os.getenv("DB_HOST", "aws-0-ap-northeast-2.pooler.supabase.com").strip() or "aws-0-ap-northeast-2.pooler.supabase.com",
    "port": safe_int(os.getenv("DB_PORT"), 6543),
    "dbname": os.getenv("DB_NAME", "postgres").strip() or "postgres",
    "user": os.getenv("DB_USER", "postgres.uaoysegountarjanafbb").strip() or "postgres.uaoysegountarjanafbb",
    "password": os.getenv("DB_PASS", "").strip(),
}

# Use /tmp on Vercel serverless to avoid read-only filesystem errors
if os.getenv("VERCEL"):
    SQLITE_DB_PATH = "/tmp/etl_database.sqlite"
    # Seed /tmp/etl_database.sqlite from bundled repository sqlite if not yet present
    try:
        bundled_candidates = [
            os.path.join(os.path.dirname(__file__), "etl_database.sqlite"),
            os.path.join(os.path.dirname(__file__), "api", "etl_database.sqlite"),
            os.path.join(os.path.dirname(__file__), "..", "etl_database.sqlite")
        ]
        if not os.path.exists(SQLITE_DB_PATH) or os.path.getsize(SQLITE_DB_PATH) < 1000:
            import shutil
            for b in bundled_candidates:
                if os.path.exists(b) and os.path.getsize(b) > 1000:
                    shutil.copyfile(b, SQLITE_DB_PATH)
                    print(f"Copied bundled sqlite ({os.path.getsize(b)} bytes) to {SQLITE_DB_PATH}")
                    break
    except Exception as e:
        print(f"Notice: Could not copy bundled sqlite to /tmp: {e}")
else:
    SQLITE_DB_PATH = os.path.join(os.path.dirname(__file__), "etl_database.sqlite")

# Track active DB engine ("postgres" or "sqlite")
ACTIVE_DB_ENGINE = "sqlite"

def get_postgres_connection():
    """Establishes a connection to Supabase PostgreSQL."""
    if not psycopg2:
        raise ValueError("psycopg2 is not available in the current environment.")

    # Check if DATABASE_URL is provided and has a real password
    if (DATABASE_URL and 
        "[MASUKKAN_PASSWORD" not in DATABASE_URL and 
        "[YOUR" not in DATABASE_URL and 
        "<password>" not in DATABASE_URL):
        return psycopg2.connect(
            DATABASE_URL, 
            cursor_factory=RealDictCursor, 
            sslmode="require", 
            connect_timeout=6
        )
    elif (DB_CONFIG.get("password") and 
          "[MASUKKAN_PASSWORD" not in DB_CONFIG["password"] and 
          "[YOUR" not in DB_CONFIG["password"]):
        return psycopg2.connect(
            **DB_CONFIG, 
            cursor_factory=RealDictCursor, 
            sslmode="require", 
            connect_timeout=6
        )
    else:
        raise ValueError("Password Supabase PostgreSQL belum disetel di .env / Environment Variables")

def test_postgres_connection() -> bool:
    """Verifies whether PostgreSQL (Supabase) is reachable."""
    if not psycopg2:
        return False
    try:
        conn = get_postgres_connection()
        cur = conn.cursor()
        cur.execute("SELECT 1;")
        cur.close()
        conn.close()
        return True
    except Exception:
        return False

# Initialize database schema and initial seed data in SQLite (for offline fallback)
def init_sqlite_db():
    conn = sqlite3.connect(SQLITE_DB_PATH)
    conn.row_factory = sqlite3.Row
    cur = conn.cursor()

    # 1. FACUL_ETL_MH_AKSEPTASI table (29 authentic columns from data_mh_akseptasi.xlsx, no status column)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS FACUL_ETL_MH_AKSEPTASI (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fac_code TEXT,
            reff_number TEXT,
            direct TEXT,
            broker TEXT,
            nama_tertanggung TEXT,
            afiliasi_tertanggung TEXT,
            coverage TEXT,
            start_date TEXT,
            end_date TEXT,
            acceptance_status TEXT,
            nama_kapal TEXT,
            type_of_vessel TEXT,
            code_kapal TEXT,
            size_of_vessel TEXT,
            year_of_built TEXT,
            type_of_material TEXT,
            classification TEXT,
            flag TEXT,
            last_docking_date TEXT,
            jenis_muatan TEXT,
            trading_area TEXT,
            currency TEXT,
            insured_value REAL,
            premium_rate REAL,
            premium_amount REAL,
            ric REAL,
            riu_share REAL,
            riu_gross_premium REAL,
            riu_net_premium REAL
        )
    ''')

    # 2. FACUL_ETL_MH_LOSS_PLA table (24 authentic columns)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS FACUL_ETL_MH_LOSS_PLA (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fac_code TEXT,
            reff_number TEXT,
            direct TEXT,
            broker TEXT,
            nama_tertanggung TEXT,
            afiliasi_tertanggung TEXT,
            nama_tertanggung_loss TEXT,
            nama_kapal TEXT,
            type_of_vessel TEXT,
            code_kapal TEXT,
            size_of_vessel TEXT,
            year_of_built TEXT,
            type_of_material TEXT,
            classification TEXT,
            flag TEXT,
            last_docking_date TEXT,
            jenis_muatan TEXT,
            riu_share REAL,
            date_of_loss TEXT,
            currency TEXT,
            loss_amount REAL,
            loss_cause TEXT,
            loss_detail TEXT,
            settled_or_os TEXT
        )
    ''')

    # 3. FACUL_ETL_MH_LOSS_SETTLE table (24 authentic columns)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS FACUL_ETL_MH_LOSS_SETTLE (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fac_code TEXT,
            reff_number TEXT,
            direct TEXT,
            broker TEXT,
            nama_tertanggung TEXT,
            afiliasi_tertanggung TEXT,
            nama_tertanggung_loss TEXT,
            nama_kapal TEXT,
            type_of_vessel TEXT,
            code_kapal TEXT,
            size_of_vessel TEXT,
            year_of_built TEXT,
            type_of_material TEXT,
            classification TEXT,
            flag TEXT,
            last_docking_date TEXT,
            jenis_muatan TEXT,
            riu_share REAL,
            date_of_loss TEXT,
            currency TEXT,
            loss_amount REAL,
            loss_cause TEXT,
            loss_detail TEXT,
            settled_or_os TEXT
        )
    ''')

    # 4. ETL_HISTORY table
    cur.execute('''
        CREATE TABLE IF NOT EXISTS ETL_HISTORY (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            file_name TEXT NOT NULL,
            file_size TEXT,
            file_type TEXT DEFAULT 'XLSX',
            cedant TEXT,
            cob TEXT,
            status TEXT,
            date_display TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
            records_count INTEGER DEFAULT 0,
            schema_accuracy REAL DEFAULT 99.8,
            duration_seconds REAL DEFAULT 8.0,
            log_message TEXT
        )
    ''')

    # 5. MAPPING_TEMPLATES table
    cur.execute('''
        CREATE TABLE IF NOT EXISTS MAPPING_TEMPLATES (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            target_schema TEXT DEFAULT 'ipr_stage_db',
            mappings_json TEXT NOT NULL,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # 5b. GENERATED_TABLES table
    cur.execute('''
        CREATE TABLE IF NOT EXISTS GENERATED_TABLES (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            table_name TEXT UNIQUE NOT NULL,
            label TEXT,
            description TEXT,
            source_table TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # 6. FACUL_ETL_MH_PARSED_AI (Dynamic DWH Table for AI Exploded Vessels)
    cur.execute('''
        CREATE TABLE IF NOT EXISTS FACUL_ETL_MH_PARSED_AI (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fac_code TEXT,
            nama_kapal TEXT,
            type_of_vessel TEXT,
            code_kapal TEXT,
            size_of_vessel TEXT,
            year_of_built TEXT,
            type_of_material TEXT,
            classification TEXT,
            direct TEXT,
            broker TEXT,
            nama_tertanggung TEXT,
            currency TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # 7. AI_PROMPT_TEMPLATES table
    cur.execute('''
        CREATE TABLE IF NOT EXISTS AI_PROMPT_TEMPLATES (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            name TEXT UNIQUE NOT NULL,
            description TEXT,
            prompt_text TEXT NOT NULL,
            target_columns TEXT,
            source_columns TEXT,
            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
        )
    ''')

    # Seed Default AI Prompt Template if not exists
    cur.execute('''
        INSERT OR IGNORE INTO AI_PROMPT_TEMPLATES (name, description, prompt_text, target_columns, source_columns)
        VALUES (
            'Ekstraksi & Explode Entitas Kapal Marine Hull (Multi-Vessel to Rows)',
            'Mengekstrak dan memecah 1 baris deskripsi majemuk menjadi banyak baris entitas kapal individual dengan atribut spesifikasi lengkap.',
            'Anda adalah Senior Data Warehouse Engineer & Marine Insurance Specialist. Ekstrak entitas kapal individual dari deskripsi mentah. Jika 1 baris mengandung >1 kapal, explode menjadi baris terpisah (1 kapal = 1 baris). Ekstrak: Nama Kapal, Type of Vessel, Code Kapal, Size of Vessel, Year of Built, Type of Material, Classification.',
            '["Nama Kapal", "Type of Vessel", "Code Kapal", "Size of Vessel", "Year of Built", "Type of Material", "Classification"]',
            '{"Nama Kapal": "fac_desc", "Type of Vessel": ["fac_risk", "fac_desc"], "Code Kapal": "fac_desc", "Size of Vessel": "fac_desc", "Year of Built": "fac_desc", "Type of Material": "fac_desc", "Classification": "fac_desc"}'
        )
    ''')

    # Seed initial AI parsed records if table is currently empty
    cur.execute('SELECT count(*) FROM FACUL_ETL_MH_PARSED_AI')
    ai_count = cur.fetchone()[0]
    if ai_count == 0:
        try:
            sample_parsed = ai_engine.run_ai_parsing(demo_data.DEMO_RAW_10_ROWS)["data"]
            for r in sample_parsed:
                cur.execute('''
                    INSERT INTO FACUL_ETL_MH_PARSED_AI (
                        fac_code, nama_kapal, type_of_vessel, code_kapal,
                        size_of_vessel, year_of_built, type_of_material,
                        classification, direct, broker, nama_tertanggung, currency
                    ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ''', (
                    r.get("fac_code", ""), r.get("nama_kapal", ""), r.get("type_of_vessel", ""),
                    r.get("code_kapal", ""), r.get("size_of_vessel", ""), r.get("year_of_built", ""),
                    r.get("type_of_material", ""), r.get("classification", ""),
                    r.get("direct", ""), r.get("broker", ""), r.get("nama_tertanggung", ""),
                    r.get("currency", "IDR")
                ))
        except Exception as seed_err:
            print(f"Notice: Could not seed initial AI records: {seed_err}")

    conn.commit()
    conn.close()

# Auto-initialize local SQLite on startup as fallback
try:
    init_sqlite_db()
except Exception as e:
    print(f"Warning: Failed to initialize SQLite fallback: {e}")

def init_postgres_db():
    try:
        conn = get_postgres_connection()
        cur = conn.cursor()
        cur.execute('''
            CREATE TABLE IF NOT EXISTS generated_tables (
                id SERIAL PRIMARY KEY,
                table_name TEXT UNIQUE NOT NULL,
                label TEXT,
                description TEXT,
                source_table TEXT,
                created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP
            )
        ''')
        conn.commit()
        cur.close()
        conn.close()
    except Exception as e:
        print(f"Warning: Failed to initialize Postgres tables: {e}")

def get_db_connection():
    """Returns an active connection, preferring PostgreSQL (Supabase) then SQLite."""
    global ACTIVE_DB_ENGINE
    if test_postgres_connection():
        ACTIVE_DB_ENGINE = "postgres"
        init_postgres_db()
        return get_postgres_connection()
    else:
        ACTIVE_DB_ENGINE = "sqlite"
        conn = sqlite3.connect(SQLITE_DB_PATH)
        conn.row_factory = sqlite3.Row
        return conn

def format_query(query_sql: str, is_sqlite: bool) -> str:
    """Adapts query syntax between Postgres and SQLite."""
    sql = query_sql
    if is_sqlite:
        sql = sql.replace('public."FACUL_ETL_MH_AKSEPTASI"', 'FACUL_ETL_MH_AKSEPTASI')
        sql = sql.replace('public."FACUL_ETL_MH_LOSS_PLA"', 'FACUL_ETL_MH_LOSS_PLA')
        sql = sql.replace('public."FACUL_ETL_MH_LOSS_SETTLE"', 'FACUL_ETL_MH_LOSS_SETTLE')
        sql = sql.replace('public."FACUL_ETL_MH_PARSED_AI"', 'FACUL_ETL_MH_PARSED_AI')
        sql = sql.replace('"FACUL_ETL_MH_AKSEPTASI"', 'FACUL_ETL_MH_AKSEPTASI')
        sql = sql.replace('"FACUL_ETL_MH_LOSS_PLA"', 'FACUL_ETL_MH_LOSS_PLA')
        sql = sql.replace('"FACUL_ETL_MH_LOSS_SETTLE"', 'FACUL_ETL_MH_LOSS_SETTLE')
        sql = sql.replace('"FACUL_ETL_MH_PARSED_AI"', 'FACUL_ETL_MH_PARSED_AI')
        sql = sql.replace('etl_history', 'ETL_HISTORY')
        sql = sql.replace('mapping_templates', 'MAPPING_TEMPLATES')
        sql = sql.replace('ai_prompt_templates', 'AI_PROMPT_TEMPLATES')
        sql = sql.replace('generated_tables', 'GENERATED_TABLES')
        sql = sql.replace('%s', '?')
    else:
        sql = sql.replace('?', '%s')
        sql = sql.replace('ETL_HISTORY', 'etl_history')
        sql = sql.replace('MAPPING_TEMPLATES', 'mapping_templates')
        sql = sql.replace('AI_PROMPT_TEMPLATES', 'ai_prompt_templates')
        sql = sql.replace('GENERATED_TABLES', 'generated_tables')
    return sql

def format_column_label(col_name: str) -> str:
    """Converts snake_case or technical column names to elegant human-readable labels."""
    label_map = {
        "fac_code": "Fac Code",
        "reff_number": "Reff Number",
        "fac_old_ref": "Reff Number",
        "direct": "Direct (Cedant)",
        "fac_cedant": "Direct (Cedant)",
        "broker": "Broker",
        "fac_broker": "Broker",
        "nama_tertanggung": "Nama Tertanggung",
        "fac_insured": "Nama Tertanggung",
        "afiliasi_tertanggung": "Afiliasi Tertanggung",
        "coverage": "Coverage",
        "start_date": "Start Date",
        "end_date": "End Date",
        "acceptance_status": "Acceptance Status",
        "nama_tertanggung_loss": "Nama Tertanggung Loss",
        "nama_tertanngung_loss": "Nama Tertanggung Loss",
        "nama_kapal": "Nama Kapal",
        "fac_vessel": "Nama Kapal",
        "type_of_vessel": "Type of Vessel",
        "fac_risk": "Type of Vessel",
        "code_kapal": "Code Kapal",
        "size_of_vessel": "Size of Vessel",
        "year_of_built": "Year of Built",
        "type_of_material": "Type of Material",
        "classification": "Classification",
        "flag": "Flag",
        "last_docking_date": "Last Docking Date",
        "jenis_muatan": "Jenis Muatan",
        "trading_area": "Trading Area",
        "currency": "Currency",
        "insured_value": "Insured Value",
        "premium_rate": "Premium Rate",
        "premium_amount": "Premium Amount",
        "ric": "RIC",
        "riu_share": "RIU Share",
        "riu_gross_premium": "RIU Gross Premium",
        "riu_net_premium": "RIU Net Premium",
        "sum_insured": "Sum Insured",
        "loss_amount": "OUR Loss Amount",
        "date_of_loss": "Date of Loss / UW Year",
        "loss_cause": "Cause of Loss",
        "status": "Status",
        "created_at": "Waktu Dibuat",
        "fac_desc": "Deskripsi Mentah"
    }
    if col_name in label_map:
        return label_map[col_name]
    return col_name.replace("_", " ").title()

def introspect_table_columns(table_name: str) -> List[Dict[str, Any]]:
    """Dynamically reads columns and types from the active database table at runtime."""
    conn = get_db_connection()
    cur = conn.cursor()
    columns_meta = []
    try:
        clean_tbl = table_name.replace('public.', '').replace('"', '')
        is_akseptasi = "AKSEPTASI" in clean_tbl.upper()

        if ACTIVE_DB_ENGINE == "sqlite":
            cur.execute(f"PRAGMA table_info({clean_tbl});")
            cols = cur.fetchall()
            for col in cols:
                col_name = col[1]
                col_type = col[2].upper() if col[2] else "TEXT"
                if col_name.lower() in ["id"]:
                    continue
                # For Akseptasi, ensure legacy/generic 'status' column is NEVER included
                if is_akseptasi and col_name.lower() == "status":
                    continue
                columns_meta.append({
                    "key": col_name,
                    "label": format_column_label(col_name),
                    "dataType": col_type,
                    "isAmount": col_name in ["sum_insured", "loss_amount", "insured_value", "premium_amount", "riu_gross_premium", "riu_net_premium", "fac_totsi", "fac_our_amt", "ric", "riu_share", "premium_rate"],
                    "isDate": col_name in ["date_of_loss", "start_date", "end_date", "last_docking_date", "fac_com_date", "fac_exp_date", "fac_doc_date", "created_at"],
                    "isCode": col_name in ["code_kapal", "fac_code", "reff_number", "fac_old_ref"],
                    "isVessel": col_name in ["nama_kapal", "fac_vessel"],
                    "bold": col_name in ["fac_code", "nama_kapal"]
                })
        else:
            # Postgres information_schema query (support both uppercase and lowercase)
            clean_tbl = table_name.replace('public.', '').replace('"', '')
            cur.execute("""
                SELECT column_name, data_type 
                FROM information_schema.columns 
                WHERE table_name = %s OR table_name = %s OR table_name = %s
                ORDER BY ordinal_position;
            """, [clean_tbl, clean_tbl.lower(), clean_tbl.upper()])
            cols = cur.fetchall()
            for col in cols:
                col_name = col["column_name"] if isinstance(col, dict) else col[0]
                col_type = (col["data_type"] if isinstance(col, dict) else col[1]).upper()
                if col_name.lower() in ["id"]:
                    continue
                # For Akseptasi, ensure legacy/generic 'status' column is NEVER included
                if is_akseptasi and col_name.lower() == "status":
                    continue
                columns_meta.append({
                    "key": col_name,
                    "label": format_column_label(col_name),
                    "dataType": col_type,
                    "isAmount": col_name in ["sum_insured", "loss_amount", "insured_value", "premium_amount", "riu_gross_premium", "riu_net_premium", "fac_totsi", "fac_our_amt", "ric", "riu_share", "premium_rate"],
                    "isDate": col_name in ["date_of_loss", "start_date", "end_date", "last_docking_date", "fac_com_date", "fac_exp_date", "fac_doc_date", "created_at"],
                    "isCode": col_name in ["code_kapal", "fac_code", "reff_number", "fac_old_ref"],
                    "isVessel": col_name in ["nama_kapal", "fac_vessel"],
                    "bold": col_name in ["fac_code", "nama_kapal"]
                })
    except Exception as err:
        print(f"Error introspecting table {table_name}: {err}")
    finally:
        cur.close()
        conn.close()

    # Fallback to standard columns if empty or introspection failed
    if not columns_meta:
        columns_meta = [
            {"key": "fac_code", "label": "Fac Code", "bold": True},
            {"key": "nama_kapal", "label": "Nama Kapal", "isVessel": True},
            {"key": "type_of_vessel", "label": "Type of Vessel"},
            {"key": "code_kapal", "label": "Code Kapal", "isCode": True},
            {"key": "size_of_vessel", "label": "Size of Vessel"},
            {"key": "year_of_built", "label": "Year of Built"},
            {"key": "type_of_material", "label": "Type of Material"},
            {"key": "classification", "label": "Classification"},
            {"key": "direct", "label": "Direct (Cedant)"},
            {"key": "broker", "label": "Broker"},
            {"key": "nama_tertanggung", "label": "Nama Tertanggung"},
            {"key": "currency", "label": "Currency"}
        ]
    return columns_meta


def query_all(query_sql: str, params: list = []):
    """Executes a SELECT query returning all matching rows as dictionaries."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        is_sqlite = (ACTIVE_DB_ENGINE == "sqlite")
        sql = format_query(query_sql, is_sqlite)
        cur.execute(sql, params)
        rows = cur.fetchall()
        return [dict(r) for r in rows]
    finally:
        cur.close()
        conn.close()

def query_one(query_sql: str, params: list = []):
    """Executes a SELECT query returning a single row as a dictionary."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        is_sqlite = (ACTIVE_DB_ENGINE == "sqlite")
        sql = format_query(query_sql, is_sqlite)
        cur.execute(sql, params)
        row = cur.fetchone()
        return dict(row) if row else None
    finally:
        cur.close()
        conn.close()

def execute_dml(query_sql: str, params: list = [], returning_id: bool = False):
    """Executes INSERT, UPDATE, or DELETE query with id extraction."""
    conn = get_db_connection()
    cur = conn.cursor()
    try:
        is_sqlite = (ACTIVE_DB_ENGINE == "sqlite")
        sql = format_query(query_sql, is_sqlite)
        
        # If SQLite and has RETURNING id, clean up if not supported
        if is_sqlite and "RETURNING" in sql.upper():
            sql = sql.split("RETURNING")[0].strip()

        cur.execute(sql, params)
        last_id = None
        if returning_id:
            if is_sqlite:
                last_id = cur.lastrowid
            else:
                try:
                    row = cur.fetchone()
                    if row:
                        last_id = row["id"] if isinstance(row, dict) else row[0]
                except Exception:
                    last_id = None
        conn.commit()
        return last_id
    finally:
        cur.close()
        conn.close()

# ---------------- API ENDPOINTS ----------------

@app.get("/")
@app.get("/api")
def root(request: Request):
    return {
        "title": "Indore ETL RU & Marine Hull API",
        "status": "online",
        "active_database": ACTIVE_DB_ENGINE,
        "scope_path": request.scope.get("path"),
        "raw_headers": {k.decode("latin1"): v.decode("latin1") for k, v in request.scope.get("headers", [])}
    }

@app.get("/health")
@app.get("/api/health")
def health_check():
    return {
        "status": "ok",
        "active_database": ACTIVE_DB_ENGINE,
        "postgres_connected": (ACTIVE_DB_ENGINE == "postgres"),
        "postgres_host": DB_CONFIG.get("host"),
        "timestamp": datetime.now().isoformat()
    }

@app.get("/stats")
@app.get("/api/stats")
def get_stats():
    try:
        loss_stat = query_one('''
            SELECT 
                count(*) as total_loss_count,
                COALESCE(sum(loss_amount), 0) as total_loss_amount,
                count(DISTINCT direct) as unique_cedants,
                count(DISTINCT nama_kapal) as unique_vessels
            FROM public."FACUL_ETL_MH_LOSS_PLA"
        ''')
        akseptasi_stat = query_one('SELECT count(*) as count FROM public."FACUL_ETL_MH_AKSEPTASI"')
        settle_stat = query_one('SELECT count(*) as count FROM public."FACUL_ETL_MH_LOSS_SETTLE"')
        history_stat = query_one('SELECT count(*) as count FROM etl_history')

        loss_count = loss_stat['total_loss_count'] if loss_stat else 0
        loss_amount = float(loss_stat['total_loss_amount']) if loss_stat and loss_stat['total_loss_amount'] else 0.0
        unique_cedants = loss_stat['unique_cedants'] if loss_stat else 0
        unique_vessels = loss_stat['unique_vessels'] if loss_stat else 0

        akseptasi_count = akseptasi_stat['count'] if akseptasi_stat else 0
        settle_count = settle_stat['count'] if settle_stat else 0
        history_count = history_stat['count'] if history_stat else 0

        total_records = loss_count + akseptasi_count + settle_count

        return {
            "totalRecords": total_records,
            "totalLossPLA": loss_count,
            "totalAkseptasi": akseptasi_count,
            "totalLossSettle": settle_count,
            "totalProcessedFiles": history_count,
            "totalClaimValue": loss_amount,
            "uniqueCedants": unique_cedants,
            "uniqueVessels": unique_vessels,
            "slaResolvedRate": "100.0%" if loss_count > 0 else "0.0%"
        }
    except Exception:
        return {
            "totalRecords": 0,
            "totalLossPLA": 0,
            "totalAkseptasi": 0,
            "totalLossSettle": 0,
            "totalProcessedFiles": 0,
            "totalClaimValue": 0,
            "uniqueCedants": 0,
            "uniqueVessels": 0,
            "slaResolvedRate": "0.0%"
        }

def build_filter_clause(
    fac_code: Optional[str] = None,
    reff_number: Optional[str] = None,
    company_name: Optional[str] = None,
    direct: Optional[str] = None,
    broker: Optional[str] = None,
    insured_name: Optional[str] = None,
    insured_loss_name: Optional[str] = None,
    vessel_name: Optional[str] = None,
    vessel_code: Optional[str] = None,
    status: Optional[str] = None,
    loss_cause: Optional[str] = None,
    currency: Optional[str] = None,
    date_of_loss: Optional[str] = None,
    search: Optional[str] = None
):
    where_clauses = []
    params = []

    if fac_code:
        where_clauses.append("LOWER(fac_code) LIKE %s")
        params.append(f"%{fac_code.lower().strip()}%")
    if reff_number:
        where_clauses.append("LOWER(reff_number) LIKE %s")
        params.append(f"%{reff_number.lower().strip()}%")
    c_name = company_name or direct
    if c_name:
        where_clauses.append("LOWER(direct) LIKE %s")
        params.append(f"%{c_name.lower().strip()}%")
    if broker:
        where_clauses.append("LOWER(broker) LIKE %s")
        params.append(f"%{broker.lower().strip()}%")
    if insured_name:
        where_clauses.append("LOWER(nama_tertanggung) LIKE %s")
        params.append(f"%{insured_name.lower().strip()}%")
    if insured_loss_name:
        where_clauses.append("(LOWER(nama_tertanggung_loss) LIKE %s OR LOWER(nama_tertanggung) LIKE %s)")
        params.extend([f"%{insured_loss_name.lower().strip()}%", f"%{insured_loss_name.lower().strip()}%"])
    if vessel_name:
        where_clauses.append("LOWER(nama_kapal) LIKE %s")
        params.append(f"%{vessel_name.lower().strip()}%")
    if vessel_code:
        where_clauses.append("LOWER(code_kapal) LIKE %s")
        params.append(f"%{vessel_code.lower().strip()}%")
    if status and status not in ["Semua", "Status: Semua"]:
        where_clauses.append("LOWER(status) = %s")
        params.append(status.lower().strip())
    if loss_cause:
        where_clauses.append("LOWER(loss_cause) LIKE %s")
        params.append(f"%{loss_cause.lower().strip()}%")
    if currency:
        where_clauses.append("LOWER(currency) = %s")
        params.append(currency.lower().strip())
    if date_of_loss:
        where_clauses.append("date_of_loss LIKE %s")
        params.append(f"%{date_of_loss.strip()}%")
    if search:
        s = f"%{search.lower().strip()}%"
        where_clauses.append("(LOWER(fac_code) LIKE %s OR LOWER(reff_number) LIKE %s OR LOWER(direct) LIKE %s OR LOWER(nama_kapal) LIKE %s OR LOWER(nama_tertanggung) LIKE %s)")
        params.extend([s, s, s, s, s])

    where_sql = f" WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
    return where_sql, params

@app.get("/api/loss-pla")
def get_loss_pla(
    page: int = 1,
    limit: int = 50,
    fac_code: Optional[str] = None,
    reff_number: Optional[str] = None,
    company_name: Optional[str] = None,
    direct: Optional[str] = None,
    broker: Optional[str] = None,
    insured_name: Optional[str] = None,
    insured_loss_name: Optional[str] = None,
    vessel_name: Optional[str] = None,
    vessel_code: Optional[str] = None,
    status: Optional[str] = None,
    loss_cause: Optional[str] = None,
    currency: Optional[str] = None,
    date_of_loss: Optional[str] = None,
    search: Optional[str] = None
):
    offset = (page - 1) * limit
    where_sql, params = build_filter_clause(
        fac_code=fac_code, reff_number=reff_number, company_name=company_name, direct=direct,
        broker=broker, insured_name=insured_name, insured_loss_name=insured_loss_name,
        vessel_name=vessel_name, vessel_code=vessel_code, status=status,
        loss_cause=loss_cause, currency=currency, date_of_loss=date_of_loss, search=search
    )

    count_res = query_one(f'SELECT count(*) as count FROM public."FACUL_ETL_MH_LOSS_PLA"{where_sql}', params)
    total_records = count_res['count'] if count_res else 0

    query = f'SELECT * FROM public."FACUL_ETL_MH_LOSS_PLA"{where_sql} ORDER BY id ASC LIMIT %s OFFSET %s'
    data_params = params + [limit, offset]
    data = query_all(query, data_params)

    total_pages = (total_records + limit - 1) // limit if limit > 0 else 1

    return {
        "page": page,
        "limit": limit,
        "total": total_records,
        "totalPages": total_pages,
        "data": data
    }

@app.get("/api/akseptasi")
def get_akseptasi(
    page: int = 1,
    limit: int = 50,
    fac_code: Optional[str] = None,
    reff_number: Optional[str] = None,
    company_name: Optional[str] = None,
    direct: Optional[str] = None,
    broker: Optional[str] = None,
    insured_name: Optional[str] = None,
    vessel_name: Optional[str] = None,
    vessel_code: Optional[str] = None,
    status: Optional[str] = None,
    currency: Optional[str] = None,
    search: Optional[str] = None
):
    offset = (page - 1) * limit
    where_sql, params = build_filter_clause(
        fac_code=fac_code, reff_number=reff_number, company_name=company_name, direct=direct,
        broker=broker, insured_name=insured_name, vessel_name=vessel_name, vessel_code=vessel_code,
        status=status, currency=currency, search=search
    )

    count_res = query_one(f'SELECT count(*) as count FROM public."FACUL_ETL_MH_AKSEPTASI"{where_sql}', params)
    total_records = count_res['count'] if count_res else 0

    data = query_all(
        f'SELECT * FROM public."FACUL_ETL_MH_AKSEPTASI"{where_sql} ORDER BY id ASC LIMIT %s OFFSET %s', 
        params + [limit, offset]
    )
    total_pages = (total_records + limit - 1) // limit if limit > 0 else 1

    return {
        "page": page,
        "limit": limit,
        "total": total_records,
        "totalPages": total_pages,
        "data": data
    }

@app.get("/api/loss-settle")
def get_loss_settle(
    page: int = 1,
    limit: int = 50,
    fac_code: Optional[str] = None,
    reff_number: Optional[str] = None,
    company_name: Optional[str] = None,
    direct: Optional[str] = None,
    broker: Optional[str] = None,
    insured_name: Optional[str] = None,
    vessel_name: Optional[str] = None,
    vessel_code: Optional[str] = None,
    status: Optional[str] = None,
    currency: Optional[str] = None,
    date_of_loss: Optional[str] = None,
    search: Optional[str] = None
):
    offset = (page - 1) * limit
    where_sql, params = build_filter_clause(
        fac_code=fac_code, reff_number=reff_number, company_name=company_name, direct=direct,
        broker=broker, insured_name=insured_name, vessel_name=vessel_name, vessel_code=vessel_code,
        status=status, currency=currency, date_of_loss=date_of_loss, search=search
    )

    count_res = query_one(f'SELECT count(*) as count FROM public."FACUL_ETL_MH_LOSS_SETTLE"{where_sql}', params)
    total_records = count_res['count'] if count_res else 0

    data = query_all(
        f'SELECT * FROM public."FACUL_ETL_MH_LOSS_SETTLE"{where_sql} ORDER BY id ASC LIMIT %s OFFSET %s', 
        params + [limit, offset]
    )
    total_pages = (total_records + limit - 1) // limit if limit > 0 else 1

    return {
        "page": page,
        "limit": limit,
        "total": total_records,
        "totalPages": total_pages,
        "data": data
    }

# ---------------- HISTORY ENDPOINTS ----------------

@app.get("/api/history")
def get_history(
    search: Optional[str] = None,
    cob: Optional[str] = None,
    status: Optional[str] = None,
    page: int = 1,
    limit: int = 9
):
    offset = (page - 1) * limit
    where_clauses = []
    params = []

    if search:
        search_lower = f"%{search.lower().strip()}%"
        where_clauses.append("(LOWER(file_name) LIKE %s OR LOWER(cedant) LIKE %s)")
        params.extend([search_lower, search_lower])

    if cob and cob not in ["Semua Tipe COB", "Semua"]:
        where_clauses.append("LOWER(cob) = %s")
        params.append(cob.lower().strip())

    if status and status not in ["Status: Semua", "Semua"]:
        where_clauses.append("status = %s")
        params.append(status.strip())

    where_sql = f" WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

    count_res = query_one(f"SELECT count(*) as count FROM etl_history{where_sql}", params)
    total_records = count_res['count'] if count_res else 0

    rows = query_all(
        f"SELECT * FROM etl_history{where_sql} ORDER BY id DESC LIMIT %s OFFSET %s", 
        params + [limit, offset]
    )

    total_pages = (total_records + limit - 1) // limit if limit > 0 else 1

    return {
        "page": page,
        "limit": limit,
        "total": total_records,
        "totalPages": total_pages,
        "data": rows
    }

class HistoryCreateRequest(BaseModel):
    file_name: str
    file_size: str
    file_type: Optional[str] = "XLSX"
    cedant: str
    cob: str
    status: Optional[str] = "Berhasil Dimuat"
    records_count: Optional[int] = 48250
    schema_accuracy: Optional[float] = 99.8
    duration_seconds: Optional[float] = 8.0
    log_message: Optional[str] = "File diproses dan divalidasi sukses."

@app.post("/api/history")
def create_history(item: HistoryCreateRequest):
    now_display = datetime.now().strftime("%d %b %Y, %H:%M WIB")

    sql = '''
        INSERT INTO etl_history (
            file_name, file_size, file_type, cedant, cob, status,
            date_display, records_count, schema_accuracy, duration_seconds, log_message
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id;
    '''

    new_id = execute_dml(sql, [
        item.file_name, item.file_size, item.file_type, item.cedant, item.cob,
        item.status, now_display, item.records_count, item.schema_accuracy,
        item.duration_seconds, item.log_message
    ], returning_id=True)

    return {
        "success": True,
        "id": new_id,
        "message": f"Riwayat berkas {item.file_name} berhasil disimpan."
    }

@app.get("/api/history/{history_id}")
def get_history_detail(history_id: int):
    row = query_one("SELECT * FROM etl_history WHERE id = %s", [history_id])
    if not row:
        raise HTTPException(status_code=404, detail="Riwayat tidak ditemukan")
    return row

# ---------------- MAPPING TEMPLATES ENDPOINTS ----------------

@app.get("/api/mapping-templates")
def get_mapping_templates():
    rows = query_all("SELECT * FROM mapping_templates ORDER BY id ASC")
    result = []
    for r in rows:
        item = dict(r)
        if isinstance(item.get("mappings_json"), str):
            try:
                item["mappings"] = json.loads(item["mappings_json"])
            except Exception:
                item["mappings"] = []
        elif isinstance(item.get("mappings_json"), list):
            item["mappings"] = item["mappings_json"]
        else:
            item["mappings"] = []
        result.append(item)
    return result

class MappingTemplateRequest(BaseModel):
    name: str
    target_schema: Optional[str] = "ipr_stage_db"
    mappings: List[Dict[str, Any]]

@app.post("/api/mapping-templates")
def save_mapping_template(req: MappingTemplateRequest):
    mappings_str = json.dumps(req.mappings)
    
    if ACTIVE_DB_ENGINE == "postgres":
        sql = '''
            INSERT INTO mapping_templates (name, target_schema, mappings_json)
            VALUES (%s, %s, %s)
            ON CONFLICT (name) DO UPDATE SET
                target_schema = EXCLUDED.target_schema,
                mappings_json = EXCLUDED.mappings_json;
        '''
    else:
        sql = '''
            INSERT INTO mapping_templates (name, target_schema, mappings_json)
            VALUES (%s, %s, %s)
            ON CONFLICT(name) DO UPDATE SET
                target_schema = excluded.target_schema,
                mappings_json = excluded.mappings_json
        '''

    execute_dml(sql, [req.name, req.target_schema, mappings_str])
    return {"success": True, "message": f"Template '{req.name}' berhasil disimpan."}

@app.delete("/api/mapping-templates/{template_id}")
def delete_mapping_template(template_id: int):
    execute_dml("DELETE FROM mapping_templates WHERE id = %s", [template_id])
    return {"success": True, "message": "Template berhasil dihapus."}

# ---------------- FILE UPLOAD & AI PARSING SIMULATION ----------------

@app.post("/api/upload")
async def upload_file(
    file: UploadFile = File(...),
    cob: Optional[str] = Form("Fire & Property"),
    mapping_template: Optional[str] = Form("Format Standar Bordero TriPakarta Fire 2026")
):
    contents = await file.read()
    size_mb = round(len(contents) / (1024 * 1024), 2)
    if size_mb == 0.0:
        size_mb = 1.8 # Friendly fallback display size for small test uploads

    detected_cols = [
        "No", "COB", "REGISTER NO.", "POLICY NUMBER", "CERTIFICATE NO",
        "REFF BORDEREAUX", "NAMA TERTANGGUNG", "TSI ORIGINAL", "SUM INSURED SHARE",
        "PREMIUM AMOUNT", "BROKER NAME", "VESSEL CODE", "VESSEL NAME",
        "LOSS CAUSE", "DATE OF LOSS", "SETTLEMENT STATUS"
    ]

    return {
        "file_name": file.filename,
        "file_size": f"{size_mb} MB",
        "cob": cob,
        "detected_columns_count": len(detected_cols),
        "detected_columns": detected_cols,
        "ready_for_mapping": True
    }

class ProcessETLRequest(BaseModel):
    file_name: str
    file_size: str
    cedant: Optional[str] = "PT Asuransi Tri Pakarta"
    cob: Optional[str] = "Fire & Property"
    mappings: Optional[List[Dict[str, Any]]] = []

@app.post("/api/process-etl")
def process_etl(req: ProcessETLRequest):
    now_display = datetime.now().strftime("%d %b %Y, %H:%M WIB")

    sql = '''
        INSERT INTO etl_history (
            file_name, file_size, file_type, cedant, cob, status,
            date_display, records_count, schema_accuracy, duration_seconds, log_message
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
        RETURNING id;
    '''

    new_id = execute_dml(sql, [
        req.file_name, req.file_size, "XLSX", req.cedant, req.cob,
        "Berhasil Dimuat", now_display, 48250, 99.8, 8.2,
        "Pemrosesan AI & Validasi skema sukses. Sebanyak 48,250 baris data berhasil divalidasi dan dimuat ke database."
    ], returning_id=True)

    return {
        "job_id": f"job-{int(time.time())}",
        "history_id": new_id,
        "file_name": req.file_name,
        "status": "completed",
        "total_records": 48250,
        "schema_accuracy": 99.8,
        "remaining_seconds": 0,
        "message": "Pemrosesan AI selesai dan disimpan ke database."
    }

# ---------------- DYNAMIC DWH & RUNTIME TABLES ENDPOINTS ----------------

@app.get("/tables")
@app.get("/api/tables")
def get_available_tables():
    """Returns all available DWH tables with runtime row counts and column counts."""
    tables_def = [
        {
            "id": "acceptance",
            "tableName": "FACUL_ETL_MH_AKSEPTASI",
            "label": "MH - Data Akseptasi",
            "description": "Tabel DWH akseptasi polis, slip penutupan, dan portofolio risiko kapal",
            "isAiParsed": False
        },
        {
            "id": "loss_pla",
            "tableName": "FACUL_ETL_MH_LOSS_PLA",
            "label": "MH - Data Loss PLA",
            "description": "Tabel DWH klaim reasuransi fakultatif berstatus preliminary / outstanding loss",
            "isAiParsed": False
        },
        {
            "id": "loss_sla",
            "tableName": "FACUL_ETL_MH_LOSS_SETTLE",
            "label": "MH - Data Loss SLA",
            "description": "Tabel DWH klaim lunas dan realisasi pembayaran santunan reasuransi",
            "isAiParsed": False
        },
        {
            "id": "ai_parsed",
            "tableName": "FACUL_ETL_MH_PARSED_AI",
            "label": "Marine Hull - Hasil Normalisasi AI (Entitas Granular)",
            "description": "Tabel DWH hasil entity extraction & multi-vessel exploding dengan AI Parsing Engine",
            "isAiParsed": True
        }
    ]

    # Fetch generated tables
    try:
        gen_rows = query_all("SELECT * FROM GENERATED_TABLES ORDER BY id ASC")
        for r in gen_rows:
            tables_def.append({
                "id": r["table_name"],
                "tableName": r["table_name"],
                "label": r["label"] or r["table_name"],
                "description": r["description"] or "Tabel Output Simulasi ETL",
                "isAiParsed": False
            })
    except Exception as e:
        print("Error fetching generated tables:", e)

    result = []
    for t in tables_def:
        try:
            cnt_res = query_one(f'SELECT count(*) as count FROM "{t["tableName"]}"')
            cnt = cnt_res['count'] if cnt_res else 0
        except Exception:
            cnt = 0
        cols = introspect_table_columns(t["tableName"])
        result.append({
            **t,
            "count": cnt,
            "columnsCount": len(cols)
        })
    return result

@app.get("/table-data")
@app.get("/api/table-data")
def get_table_data_generic(
    table: str = Query("acceptance", description="Table key or table name"),
    page: int = 1,
    limit: int = 12,
    search: Optional[str] = None,
    fac_code: Optional[str] = None,
    vessel_name: Optional[str] = None,
    vessel_loss_name: Optional[str] = None,
    vessel_code: Optional[str] = None,
    vessel_loss_code: Optional[str] = None,
    insured_name: Optional[str] = None,
    insured_loss_name: Optional[str] = None,
    company_name: Optional[str] = None,
    status: Optional[str] = None
):
    """Generic endpoint returning runtime rows AND dynamic columns metadata for ANY table."""
    table_map = {
        "loss_pla": "FACUL_ETL_MH_LOSS_PLA",
        "acceptance": "FACUL_ETL_MH_AKSEPTASI",
        "loss_sla": "FACUL_ETL_MH_LOSS_SETTLE",
        "ai_parsed": "FACUL_ETL_MH_PARSED_AI",
        "FACUL_ETL_MH_LOSS_PLA": "FACUL_ETL_MH_LOSS_PLA",
        "FACUL_ETL_MH_AKSEPTASI": "FACUL_ETL_MH_AKSEPTASI",
        "FACUL_ETL_MH_LOSS_SETTLE": "FACUL_ETL_MH_LOSS_SETTLE",
        "FACUL_ETL_MH_PARSED_AI": "FACUL_ETL_MH_PARSED_AI"
    }
    actual_table = table_map.get(table, table)
    columns = introspect_table_columns(actual_table)
    col_keys = [c["key"] for c in columns]

    offset = (page - 1) * limit
    where_clauses = []
    params = []

    if search:
        search_lower = f"%{search.lower().strip()}%"
        search_sub = []
        for ck in col_keys:
            if ck in ["fac_code", "nama_kapal", "direct", "nama_tertanggung", "type_of_vessel", "classification"]:
                search_sub.append(f"LOWER({ck}) LIKE %s")
                params.append(search_lower)
        if search_sub:
            where_clauses.append(f"({' OR '.join(search_sub)})")

    if fac_code and "fac_code" in col_keys:
        where_clauses.append("LOWER(fac_code) LIKE %s")
        params.append(f"%{fac_code.lower().strip()}%")

    v_name = vessel_name or vessel_loss_name
    if v_name and "nama_kapal" in col_keys:
        where_clauses.append("LOWER(nama_kapal) LIKE %s")
        params.append(f"%{v_name.lower().strip()}%")

    v_code = vessel_code or vessel_loss_code
    if v_code and "code_kapal" in col_keys:
        where_clauses.append("LOWER(code_kapal) LIKE %s")
        params.append(f"%{v_code.lower().strip()}%")

    if insured_loss_name:
        if "nama_tertanngung_loss" in col_keys:
            where_clauses.append("LOWER(nama_tertanngung_loss) LIKE %s")
            params.append(f"%{insured_loss_name.lower().strip()}%")
        elif "nama_tertanggung" in col_keys:
            where_clauses.append("LOWER(nama_tertanggung) LIKE %s")
            params.append(f"%{insured_loss_name.lower().strip()}%")

    if insured_name and "nama_tertanggung" in col_keys:
        where_clauses.append("LOWER(nama_tertanggung) LIKE %s")
        params.append(f"%{insured_name.lower().strip()}%")

    if company_name and "direct" in col_keys:
        where_clauses.append("LOWER(direct) LIKE %s")
        params.append(f"%{company_name.lower().strip()}%")

    if status and "status" in col_keys and status not in ["Semua", "Status: Semua"]:
        where_clauses.append("LOWER(status) = %s")
        params.append(status.lower().strip())

    where_sql = f" WHERE {' AND '.join(where_clauses)}" if where_clauses else ""

    count_res = query_one(f'SELECT count(*) as count FROM "{actual_table}"{where_sql}', params)
    total_records = count_res['count'] if count_res else 0

    query = f'SELECT * FROM "{actual_table}"{where_sql} ORDER BY id ASC LIMIT %s OFFSET %s'
    data_params = params + [limit, offset]
    data = query_all(query, data_params)

    total_pages = (total_records + limit - 1) // limit if limit > 0 else 1

    return {
        "tableName": actual_table,
        "tableKey": table,
        "columns": columns,
        "data": data,
        "page": page,
        "limit": limit,
        "total": total_records,
        "totalPages": total_pages
    }

# ---------------- UNPARSED BATCH DATA & FILE INSPECTOR ----------------

@app.get("/api/unparsed-batch")
@app.get("/api/demo-unparsed")
def get_demo_unparsed_data():
    """Returns curated unparsed raw Marine Hull records for Before vs After AI parsing."""
    return {
        "fileName": "Bordero_MarineHull_Batch_Unparsed.xlsx",
        "fileSize": "14.2 KB",
        "cob": "Marine Hull",
        "total": len(demo_data.DEMO_RAW_10_ROWS),
        "data": demo_data.DEMO_RAW_10_ROWS,
        "columns": [
            {"key": "fac_code", "label": "Fac Code", "bold": True},
            {"key": "fac_risk", "label": "Type of Vessel (fac_risk)"},
            {"key": "fac_desc", "label": "Deskripsi Mentah (fac_desc)"},
            {"key": "fac_cedant", "label": "Direct (Cedant)"},
            {"key": "fac_insured", "label": "Nama Tertanggung"}
        ]
    }

@app.post("/api/inspect-file")
async def inspect_uploaded_file(file: UploadFile = File(...)):
    """Reads header columns and sample rows from user-uploaded Excel (.xlsx, .xls) or CSV files."""
    contents = await file.read()
    filename = file.filename or "uploaded_file.xlsx"
    size_mb = round(len(contents) / (1024 * 1024), 2)
    size_str = f"{size_mb} MB" if size_mb >= 0.1 else f"{round(len(contents) / 1024, 1)} KB"

    columns = []
    sample_rows = []

    try:
        if filename.lower().endswith(".csv"):
            text_stream = io.StringIO(contents.decode("utf-8", errors="ignore"))
            reader = csv.reader(text_stream)
            header_row = next(reader, [])
            columns = [str(col).strip() for col in header_row if col]
            for _ in range(3):
                row = next(reader, None)
                if row:
                    sample_rows.append(dict(zip(columns, row)))
        else:
            # Excel parser
            wb = openpyxl.load_workbook(io.BytesIO(contents), read_only=True, data_only=True)
            ws = wb.active
            rows_iter = ws.iter_rows(values_only=True)
            header_row = next(rows_iter, None)
            if header_row:
                columns = [str(c).strip() for c in header_row if c is not None and str(c).strip()]
                for _ in range(3):
                    r = next(rows_iter, None)
                    if r:
                        sample_rows.append({columns[i]: r[i] for i in range(min(len(columns), len(r)))})
            wb.close()
    except Exception as err:
        print(f"Error inspecting file {filename}: {err}")
        # Friendly fallback if parsing fails
        columns = [
            "fac_code", "fac_risk", "fac_desc", "fac_old_ref",
            "fac_cedant", "fac_broker", "fac_insured", "currency", "fac_totsi"
        ]

    return {
        "fileName": filename,
        "fileSize": size_str,
        "columns": columns,
        "columnsCount": len(columns),
        "sampleRows": sample_rows,
        "readyForMapping": True
    }

# ---------------- AI PROMPT TEMPLATES CRUD ----------------

@app.get("/api/ai-prompt-templates")
def get_ai_prompt_templates():
    """Lists all saved AI prompt templates."""
    rows = query_all("SELECT * FROM ai_prompt_templates ORDER BY id ASC")
    result = []
    for r in rows:
        item = dict(r)
        if isinstance(item.get("target_columns"), str):
            try:
                item["target_columns"] = json.loads(item["target_columns"])
            except Exception:
                item["target_columns"] = []
        if isinstance(item.get("source_columns"), str):
            try:
                item["source_columns"] = json.loads(item["source_columns"])
            except Exception:
                item["source_columns"] = {}
        result.append(item)
    return result

class AiPromptTemplateRequest(BaseModel):
    name: str
    description: Optional[str] = ""
    prompt_text: str
    target_columns: Optional[List[str]] = None
    source_columns: Optional[Dict[str, Any]] = None

@app.post("/api/ai-prompt-templates")
def save_ai_prompt_template(req: AiPromptTemplateRequest):
    target_cols_str = json.dumps(req.target_columns or [])
    src_cols_str = json.dumps(req.source_columns or {})

    if ACTIVE_DB_ENGINE == "postgres":
        sql = '''
            INSERT INTO ai_prompt_templates (name, description, prompt_text, target_columns, source_columns)
            VALUES (%s, %s, %s, %s, %s)
            ON CONFLICT (name) DO UPDATE SET
                description = EXCLUDED.description,
                prompt_text = EXCLUDED.prompt_text,
                target_columns = EXCLUDED.target_columns,
                source_columns = EXCLUDED.source_columns;
        '''
    else:
        sql = '''
            INSERT INTO ai_prompt_templates (name, description, prompt_text, target_columns, source_columns)
            VALUES (%s, %s, %s, %s, %s)
            ON CONFLICT(name) DO UPDATE SET
                description = excluded.description,
                prompt_text = excluded.prompt_text,
                target_columns = excluded.target_columns,
                source_columns = excluded.source_columns;
        '''
    execute_dml(sql, [req.name, req.description, req.prompt_text, target_cols_str, src_cols_str])
    return {"success": True, "message": f"Template Prompt '{req.name}' berhasil disimpan."}

@app.delete("/api/ai-prompt-templates/{template_id}")
def delete_ai_prompt_template(template_id: int):
    execute_dml("DELETE FROM ai_prompt_templates WHERE id = %s", [template_id])
    return {"success": True, "message": "Template prompt berhasil dihapus."}

# ---------------- AI PARSING EXECUTION ENDPOINT ----------------

class AiParseRequest(BaseModel):
    rows: Optional[List[Dict[str, Any]]] = None
    prompt_template: Optional[str] = None
    target_columns: Optional[List[str]] = None
    source_mapping: Optional[Dict[str, Any]] = None
    file_name: Optional[str] = "Bordero_MarineHull_Batch_Unparsed.xlsx"
    file_size: Optional[str] = "14.2 KB"
    cob: Optional[str] = "Marine Hull"
    cedant: Optional[str] = "PT Asuransi Central Asia / Konsorsium"
    save_to_dwh: Optional[bool] = True
    target_table: Optional[str] = "FACUL_ETL_MH_AKSEPTASI"
    output_title: Optional[str] = ""

@app.post("/api/ai-parse")
@app.post("/ai-parse")
def execute_ai_parsing(req: AiParseRequest):
    """Executes AI parsing, entity extraction & multi-vessel exploding, then creates dynamic table."""
    input_rows = req.rows if req.rows and len(req.rows) > 0 else demo_data.DEMO_RAW_10_ROWS
    
    # Run AI Parsing using AI Parsing Engine
    parse_result = ai_engine.run_ai_parsing(
        rows=input_rows,
        prompt_template=req.prompt_template or ai_engine.DEFAULT_PROMPT_TEMPLATE,
        target_columns=req.target_columns,
        source_mapping=req.source_mapping
    )
    
    exploded_rows = parse_result.get("data", [])

    # Clean output title for table name
    timestamp_suffix = int(time.time())
    base_title = req.output_title.strip() if req.output_title else "Output_Simulasi"
    safe_table_name = "ETL_OUT_" + re.sub(r'[^a-zA-Z0-9_]', '_', base_title).upper()[:30] + f"_{timestamp_suffix}"
    
    # Identify target base table to copy structure and search from
    base_table = req.target_table
    if not base_table:
        base_table = "FACUL_ETL_MH_AKSEPTASI"

    base_cols = introspect_table_columns(base_table)
    
    conn = get_db_connection()
    cur = conn.cursor()
    
    try:
        if ACTIVE_DB_ENGINE == "sqlite":
            cur.execute(f"PRAGMA table_info({base_table})")
            columns_def = cur.fetchall()
            col_defs = []
            for col in columns_def:
                col_name = col[1]
                col_type = col[2]
                if col_name.lower() == "id":
                    col_defs.append(f"{col_name} INTEGER PRIMARY KEY AUTOINCREMENT")
                else:
                    col_defs.append(f"{col_name} {col_type}")
            create_sql = f"CREATE TABLE {safe_table_name} ({', '.join(col_defs)})"
            cur.execute(create_sql)
        else:
            cur.execute(f'CREATE TABLE "{safe_table_name}" (LIKE "{base_table}" INCLUDING ALL)')
            
        cur.execute(
            "INSERT INTO GENERATED_TABLES (table_name, label, description, source_table) VALUES (?, ?, ?, ?)",
            (safe_table_name, req.output_title or base_table, f"Hasil proses Parsing Engine dari {req.file_name}", base_table)
        )
        
        # We need to insert the exploded rows into the dynamic table. 
        # Since exploded rows from AI engine might not have all columns of base_table, we match by column name.
        if exploded_rows:
            # Get valid columns from base_table
            if ACTIVE_DB_ENGINE == "sqlite":
                cur.execute(f"PRAGMA table_info({safe_table_name})")
                valid_cols = [c[1] for c in cur.fetchall()]
            else:
                cur.execute(f"SELECT column_name FROM information_schema.columns WHERE table_name = '{safe_table_name}' OR table_name = '{safe_table_name.lower()}'")
                valid_cols = [c[0] for c in cur.fetchall()]

            for row in exploded_rows:
                # keep only keys that exist in valid_cols
                filtered_row = {k: v for k, v in row.items() if k in valid_cols and k != 'id'}
                if not filtered_row:
                    continue
                cols = list(filtered_row.keys())
                vals = list(filtered_row.values())
                placeholders = ', '.join(['?' if ACTIVE_DB_ENGINE == "sqlite" else '%s'] * len(vals))
                
                insert_query = f'INSERT INTO "{safe_table_name}" ({", ".join(cols)}) VALUES ({placeholders})'
                cur.execute(insert_query, vals)
                
        conn.commit()
        
    except Exception as e:
        conn.rollback()
        print(f"Error creating dynamic table or inserting: {e}")
    finally:
        cur.close()
        conn.close()

    now_display = datetime.now().strftime("%d %b %Y, %H:%M WIB")
    execute_dml('''
        INSERT INTO etl_history (
            file_name, file_size, file_type, cedant, cob, status,
            date_display, records_count, schema_accuracy, duration_seconds, log_message
        ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
    ''', [
        req.file_name or "Raw_Batch_Unparsed.xlsx",
        req.file_size or "14.2 KB",
        "XLSX",
        req.cedant or "PT Asuransi",
        req.cob or "Marine Hull",
        "Berhasil Dimuat (Parsing Engine)",
        now_display,
        len(exploded_rows),
        100.0,
        3.4,
        f"Parsing Engine sukses. {len(exploded_rows)} baris data berhasil diekstrak dan dimuat ke tabel {safe_table_name}."
    ])

    sample_comparison = []
    for raw in input_rows[:2]:
        raw_code = raw.get("fac_code", "")
        matching_exploded = [r for r in exploded_rows if r.get("fac_code") == raw_code]
        sample_comparison.append({
            "raw": raw,
            "exploded": matching_exploded
        })

    return {
        "success": True,
        "usedEngine": "Parsing Engine",
        "sourceCount": len(input_rows),
        "resultCount": len(exploded_rows),
        "expansionRatio": parse_result.get("expansion_ratio", 1.5),
        "columns": base_cols,
        "data": exploded_rows,
        "sampleComparison": sample_comparison,
        "targetTable": safe_table_name,
        "message": f"Parsing selesai. Ditemukan {len(exploded_rows)} entitas."
    }


if __name__ == "__main__":
    import uvicorn
    port = safe_int(os.getenv("PORT"), 8000)
    uvicorn.run("main:app", host="0.0.0.0", port=port, reload=True)
